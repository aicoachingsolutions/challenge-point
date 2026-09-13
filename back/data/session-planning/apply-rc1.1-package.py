"""
Session Planning Model RC1.1 — apply Christian's RC1.1 package without losing Cycle 8.

Christian, 2026-09-13: the attached Session Planning Model RC1.1 is canonical. It adds two Guided
Learning Goals (A05 Progress the Attack, A06 Finish Attacks), their entry phrases and Engine
Translation rows, and an RPC Routing sheet from the routing audit.

THE RECEIVED FILE WAS BUILT FROM THE ORIGINAL RC1 WORKBOOK, not from the canonical one here, which
already carries his Cycle 8 Canonical Decisions (apply-cycle8-decisions.py):
  * Entry Language — the 57 phrases Cycle 8 approved are absent (19 rows received; 69 before RC1.1).
  * Engine Translation — the nine approved Primary Game Problem mappings read TBD again, and so do
    the two intentional gaps, A01 and A04, which Cycle 8 recorded as EMPTY because "No placeholder
    mappings should be introduced".
His covering email withdraws none of those decisions, so this is the base file's age, not a
decision. Both sets are kept:

    base     = the received RC1.1 workbook, as sent: his structure, new rows and new sheet
    restored = every Cycle 8 Entry Language phrase and Engine Translation value the base lacks

FAIL LOUDLY RATHER THAN CHOOSE. The script stops if the received workbook changes anything Cycle 8
also decided (a phrase routed to a different goal, a populated translation that differs) or edits
an existing Learning Goal, Practice Situation or Governance rule. Those would be real decisions,
and picking one silently is how authored knowledge has been lost here before.

THE CYCLE 8 DECISIONS ARE READ FROM GIT, never from the canonical file this script writes. The first
version read them from that file; one faulty run (openpyxl's cell() silently ignores None, so the
intentional EMPTY gaps stayed TBD) was enough for the re-run to read the fault back as "Cycle 8" and
faithfully restore it. Pinning the source to the commit before RC1.1 makes a re-run reproduce the
same workbook no matter what the previous run wrote.

Run from the repo root, then project:
    python back/data/session-planning/apply-rc1.1-package.py
    python back/data/session-planning/project-workbook.py
"""

import hashlib
import io
import json
import pathlib
import subprocess
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

HERE = pathlib.Path(__file__).parent
RECEIVED = HERE / "received" / "session-planning-model.rc1.1.received.xlsx"
CANONICAL = HERE / "session-planning-model.rc1.xlsx"
REPORT = HERE / "rc1.1-package-application.json"

# The canonical workbook as it stood with Cycle 8 applied and before RC1.1: the commit that added the
# received RC1.1 files, which did not touch the canonical workbook.
CYCLE8_REF = "33b21a7"
CYCLE8_PATH = "back/data/session-planning/session-planning-model.rc1.xlsx"


def cycle8_blob():
    return subprocess.run(["git", "show", f"{CYCLE8_REF}:{CYCLE8_PATH}"], cwd=HERE, capture_output=True, check=True).stdout

HEADER_ROW = 1
UNPOPULATED = "TBD"


def sha256(data):
    return hashlib.sha256(data if isinstance(data, bytes) else data.read_bytes()).hexdigest()


def empty(value):
    return value is None or str(value).strip() == ""


def same(a, b):
    """Equal as authored content: display order 1 and 1.0 are the same, and so are None and ''."""
    if empty(a) and empty(b):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return float(a) == float(b)
    return str(a).strip() == str(b).strip()


def table(ws):
    headers = [c.value for c in ws[HEADER_ROW]]
    records = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        values = [ws.cell(row=r, column=i + 1).value for i in range(len(headers))]
        if all(empty(v) for v in values):
            continue
        records.append((r, {str(h): v for h, v in zip(headers, values) if h}))
    return [str(h) for h in headers if h], records


def last_row(ws):
    rows = [r for r, _ in table(ws)[1]]
    return max(rows) if rows else HEADER_ROW


def fail(message):
    raise SystemExit(f"STOPPED — {message}")


def require_unchanged(sheet, key, cycle8, received, report):
    """Every Cycle 8 row must survive unchanged; the received workbook may only add rows."""
    headers, cycle8_rows = table(cycle8[sheet])
    _, received_rows = table(received[sheet])
    by_key = {str(row[key]).strip(): row for _, row in received_rows}
    for _, row in cycle8_rows:
        ident = str(row[key]).strip()
        if ident not in by_key:
            fail(f"{sheet}: '{ident}' exists in the Cycle 8 workbook but not in the received RC1.1 workbook.")
        for column in headers:
            if not same(row.get(column), by_key[ident].get(column)):
                fail(
                    f"{sheet} '{ident}' column '{column}' differs: Cycle 8 {row.get(column)!r}, "
                    f"received {by_key[ident].get(column)!r}. That is a decision, not a restoration."
                )
    known = {str(row[key]).strip() for _, row in cycle8_rows}
    added = [str(row[key]).strip() for _, row in received_rows if str(row[key]).strip() not in known]
    report["added_by_rc1_1"][sheet] = added


def main():
    if not RECEIVED.exists():
        fail(f"received workbook not found: {RECEIVED}")
    blob = cycle8_blob()
    report = {
        "received_workbook": {"path": RECEIVED.name, "sha256": sha256(RECEIVED)},
        "cycle8_workbook": {"git_ref": CYCLE8_REF, "path": CYCLE8_PATH, "sha256": sha256(blob)},
        "added_by_rc1_1": {},
        "restored_from_cycle8": {},
    }

    cycle8 = openpyxl.load_workbook(io.BytesIO(blob))
    received = openpyxl.load_workbook(RECEIVED)

    missing_sheets = [name for name in cycle8.sheetnames if name not in received.sheetnames]
    if missing_sheets:
        fail(f"the received workbook drops sheet(s) {missing_sheets}.")
    report["added_by_rc1_1"]["sheets"] = [name for name in received.sheetnames if name not in cycle8.sheetnames]

    require_unchanged("Learning Goal Registry", "ID", cycle8, received, report)
    require_unchanged("Practice Situation Registry", "ID", cycle8, received, report)
    require_unchanged("Governance", "Rule", cycle8, received, report)

    goal_ids = {str(row["ID"]).strip() for _, row in table(received["Learning Goal Registry"])[1]}

    # Entry Language: a phrase is keyed by its text. Same phrase, same goal: already there. Same
    # phrase, different goal: stop. Absent: restore it, in Cycle 8 order, after the received rows.
    phrase_ws = received["Entry Language"]
    _, received_phrases = table(phrase_ws)
    received_goal_for = {str(row["Coach Phrase"]).strip().lower(): str(row["Learning Goal ID"]).strip() for _, row in received_phrases}
    restored_phrases = []
    next_row = last_row(phrase_ws) + 1
    for _, row in table(cycle8["Entry Language"])[1]:
        phrase = str(row["Coach Phrase"]).strip()
        goal = str(row["Learning Goal ID"]).strip()
        existing = received_goal_for.get(phrase.lower())
        if existing is not None:
            if existing != goal:
                fail(f"Entry Language '{phrase}' routes to {goal} in Cycle 8 but {existing} in RC1.1.")
            continue
        if goal not in goal_ids:
            fail(f"Entry Language '{phrase}' points at {goal}, which the RC1.1 registry does not contain.")
        phrase_ws.cell(row=next_row, column=1).value = row["Coach Phrase"]
        phrase_ws.cell(row=next_row, column=2).value = row["Learning Goal ID"]
        next_row += 1
        restored_phrases.append({"phrase": phrase, "learning_goal_id": goal})
    report["added_by_rc1_1"]["Entry Language"] = [
        {"phrase": str(row["Coach Phrase"]).strip(), "learning_goal_id": str(row["Learning Goal ID"]).strip()}
        for _, row in received_phrases
        if str(row["Coach Phrase"]).strip().lower() not in {str(r["Coach Phrase"]).strip().lower() for _, r in table(cycle8["Entry Language"])[1]}
    ]
    report["restored_from_cycle8"]["Entry Language"] = restored_phrases

    # Engine Translation: a received row that is still TBD takes the Cycle 8 decision. A populated
    # received value must already agree — otherwise it is a new decision and this stops.
    translation_ws = received["Engine Translation"]
    t_headers, received_translation = table(translation_ws)
    column_of = {h: i + 1 for i, h in enumerate(t_headers)}
    received_row_for = {str(row["Learning Goal ID"]).strip(): (r, row) for r, row in received_translation}
    restored_translation = []
    for _, decided in table(cycle8["Engine Translation"])[1]:
        goal = str(decided["Learning Goal ID"]).strip()
        if goal not in received_row_for:
            fail(f"Engine Translation for {goal} exists in Cycle 8 but not in RC1.1.")
        r, row = received_row_for[goal]
        placeholder = all(
            empty(row.get(c)) or str(row.get(c)).strip().upper() == UNPOPULATED for c in ("Primary GP IDs", "Secondary GP IDs")
        )
        if not placeholder:
            for column in ("Primary GP IDs", "Secondary GP IDs"):
                if not same(row.get(column), decided.get(column)):
                    fail(f"Engine Translation {goal} '{column}' is {row.get(column)!r} in RC1.1 but {decided.get(column)!r} in Cycle 8.")
            continue
        if all(same(row.get(c), decided.get(c)) for c in ("Primary GP IDs", "Secondary GP IDs", "Notes")):
            continue
        for column in ("Primary GP IDs", "Secondary GP IDs", "Notes"):
            # Assigned through .value, never cell(..., value=...): openpyxl's cell() ignores a value of
            # None, so an intentional EMPTY decision (A01, A04) would silently stay "TBD".
            translation_ws.cell(row=r, column=column_of[column]).value = decided.get(column)
        restored_translation.append(
            {"learning_goal_id": goal, "primary": decided.get("Primary GP IDs"), "secondary": decided.get("Secondary GP IDs"),
             "notes": decided.get("Notes"), "received_placeholder": {c: row.get(c) for c in ("Primary GP IDs", "Secondary GP IDs", "Notes")}}
        )
    report["restored_from_cycle8"]["Engine Translation"] = restored_translation

    received.save(CANONICAL)
    REPORT.write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")

    print(f"Added by RC1.1: sheets {report['added_by_rc1_1']['sheets']}; "
          f"goals {report['added_by_rc1_1']['Learning Goal Registry']}; "
          f"situations {report['added_by_rc1_1']['Practice Situation Registry']}; "
          f"phrases {len(report['added_by_rc1_1']['Entry Language'])}")
    print(f"Restored from Cycle 8: {len(restored_phrases)} entry phrases, {len(restored_translation)} Engine Translation rows "
          f"({', '.join(t['learning_goal_id'] for t in restored_translation)})")
    print(f"Wrote {CANONICAL.name} and {REPORT.name}.")


if __name__ == "__main__":
    main()
