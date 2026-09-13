"""
Apply Christian's RC1.1 follow-up decisions (2026-09-13) to the canonical Session Planning workbook.

He adopted the combined Session Planning Model RC1.1 (his RC1.1 file with Cycle 8 restored) as the
canonical source, so this edits that workbook directly. The one-off combination script and the
received file it read are retired; git history keeps both.

1. ENTRY LANGUAGE — CREATION VS CONVERSION. His semantic rule, recorded so later Entry Language
   decisions follow it rather than becoming isolated vocabulary exceptions:
       language describing the CREATION of an opportunity           -> Create Scoring Chances (A03)
       language describing the EXECUTION or CONVERSION of one that
       already exists                                                -> Finish Attacks (A06)
   Applied:
       shoot, shot, shooting, convert chance -> A06
       get a shot                            -> stays A03 ("creating access to an opportunity to shoot")
       finish, score, goal                   -> stay A03 and keep disambiguating between A03 and A06

2. RPC ROUTING COLUMN — "Primary RPC ID" becomes "Routed RPC ID". The Session Planning Model states
   which context a goal routes to, not the strength of that relationship inside the context, so the
   old name read as contradicting the three SECONDARY routes (A04, D03, TA02).

Each change asserts the state it expects before writing and accepts the state it produces: the
script is idempotent, and it refuses to run over anything else.

Run from the repo root, then project:
    python back/data/session-planning/apply-rc1.1-decisions.py
    python back/data/session-planning/project-workbook.py
"""

import pathlib
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

HERE = pathlib.Path(__file__).parent
CANONICAL = HERE / "session-planning-model.rc1.xlsx"
HEADER_ROW = 1

CONVERSION_LANGUAGE = {"shoot": "A06", "shot": "A06", "shooting": "A06", "convert chance": "A06"}
FROM_GOAL = "A03"
UNCHANGED = {"get a shot": "A03", "finish": "A03", "score": "A03", "goal": "A03"}

ROUTING_SHEET = "RPC Routing"
ROUTING_HEADER_OLD = "Primary RPC ID"
ROUTING_HEADER_NEW = "Routed RPC ID"


def fail(message):
    raise SystemExit(f"STOPPED — {message}")


def main():
    wb = openpyxl.load_workbook(CANONICAL)

    ws = wb["Entry Language"]
    headers = [c.value for c in ws[HEADER_ROW]]
    phrase_col = headers.index("Coach Phrase") + 1
    goal_col = headers.index("Learning Goal ID") + 1
    row_of = {}
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        phrase = ws.cell(row=r, column=phrase_col).value
        if phrase is None or str(phrase).strip() == "":
            continue
        key = str(phrase).strip().lower()
        if key in row_of:
            fail(f"Entry Language repeats '{key}'.")
        row_of[key] = r

    moved, already = [], []
    for phrase, goal in CONVERSION_LANGUAGE.items():
        if phrase not in row_of:
            fail(f"Entry Language has no '{phrase}'.")
        cell = ws.cell(row=row_of[phrase], column=goal_col)
        current = str(cell.value).strip()
        if current == goal:
            already.append(phrase)
        elif current == FROM_GOAL:
            cell.value = goal
            moved.append(phrase)
        else:
            fail(f"'{phrase}' routes to {current}; expected {FROM_GOAL} before the decision or {goal} after it.")

    for phrase, goal in UNCHANGED.items():
        if phrase not in row_of:
            fail(f"Entry Language has no '{phrase}'.")
        current = str(ws.cell(row=row_of[phrase], column=goal_col).value).strip()
        if current != goal:
            fail(f"'{phrase}' should stay with {goal} but routes to {current}.")

    routing = wb[ROUTING_SHEET]
    header_cells = {str(c.value).strip(): c for c in routing[HEADER_ROW] if c.value is not None}
    if ROUTING_HEADER_OLD in header_cells:
        header_cells[ROUTING_HEADER_OLD].value = ROUTING_HEADER_NEW
        renamed = True
    elif ROUTING_HEADER_NEW in header_cells:
        renamed = False
    else:
        fail(f"{ROUTING_SHEET} has neither '{ROUTING_HEADER_OLD}' nor '{ROUTING_HEADER_NEW}' in its header row.")

    wb.save(CANONICAL)
    print(f"Entry Language moved to A06: {moved or 'none'}; already there: {already or 'none'}.")
    print(f"Unchanged and confirmed: {sorted(UNCHANGED)}.")
    print(f"{ROUTING_SHEET} column: {'renamed' if renamed else 'already'} '{ROUTING_HEADER_NEW}'.")


if __name__ == "__main__":
    main()
