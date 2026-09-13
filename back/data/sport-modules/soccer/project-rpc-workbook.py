"""
Representative Performance Context workbook -> JSON projection.

Same pattern as every workbook we ingest: the .xlsx stays canonical, this produces a complete
generated projection the TypeScript loader reads. Never hand-edit the JSON — regenerate it.

Run from the repo root, after resolve-rpc-staging.py:
    python back/data/sport-modules/soccer/project-rpc-workbook.py

Sheets are DECLARED here rather than discovered, and every one is required: a revision that drops
or renames a sheet fails loudly instead of projecting an empty list the loader would read as "no
knowledge". The README sheet is narrative and is deliberately not projected — the Workbook Standard,
Principle 5: narrative explanation should never be parsed by the runtime engine.

Implementation Staging IS projected, but only so the loader can refuse to claim ACTIVE status while
entries remain unresolved. Nothing reasons from it.
"""

import json
import pathlib
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

HERE = pathlib.Path(__file__).parent
WORKBOOK = HERE / "rpc-workbook.rc1.xlsx"
OUT = pathlib.Path(__file__).parents[3] / "src/system/sport-module/rpc-library.rc1.json"

HEADER_ROW = 1
TABLE_SHEETS = {
    "Registry": "registry",
    "Statements": "statements",
    "Relationships": "relationships",
    "Properties": "properties",
    "Identity Rules": "identity_rules",
    "Context Transitions": "transitions",
    "Validation Rules": "validation_rules",
    "Controlled Vocabulary": "controlled_vocabulary",
    "Implementation Staging": "implementation_staging",
    "Metadata": "metadata_rows",
}


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped != "" else None
    return value


def read_table(worksheet):
    headers = [c.value for c in worksheet[HEADER_ROW]]
    rows = []
    for raw in worksheet.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if not any(v is not None and str(v).strip() != "" for v in raw):
            continue
        rows.append({str(h): clean(v) for h, v in zip(headers, raw) if h})
    return rows


def main():
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True)
    missing = [name for name in TABLE_SHEETS if name not in workbook.sheetnames]
    if missing:
        raise SystemExit(f"RPC workbook is missing required sheet(s): {missing}")

    projection = {"source_workbook": WORKBOOK.name, "header_row": HEADER_ROW}
    for sheet_name, key in TABLE_SHEETS.items():
        projection[key] = read_table(workbook[sheet_name])
        print(f"  {sheet_name}: {len(projection[key])} rows")

    projection["metadata"] = {row["metadata_key"]: row["metadata_value"] for row in projection["metadata_rows"]}

    with OUT.open("w", encoding="utf-8") as handle:
        json.dump(projection, handle, indent=2, ensure_ascii=False, default=str)
        handle.write("\n")
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
