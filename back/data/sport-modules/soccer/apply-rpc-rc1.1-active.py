"""
Move RPC RC1.1 from PROPOSED to ACTIVE (Christian, 2026-09-13).

His condition, verbatim: "Once those changes are reflected and the validation passes, I'm comfortable
with RPC RC1.1 moving from PROPOSED to ACTIVE."

So this script checks the structural half of that condition itself, and refuses unless it holds:
  * no Implementation Staging row is NEEDS_CANONICAL_ID (the Workbook Standard's ACTIVE rule);
  * every context scores through at least one SCORING_EVENT relationship;
  * every context carries exactly one PRIMARY_SCORING_CONDITION.

The rest of "validation passes" cannot be checked from a workbook and is recorded where it was
measured (docs/HANDOFF.md, 2026-09-13):
  * unit suites pass, including primary-scoring, which requires all 22 context x game form pairs and
    all 13 guided goals, through context-gated selection, to resolve an approved event;
  * real generation for every guided goal, read slot by slot.

Sets Metadata runtime_status and every Registry row's runtime_status to ACTIVE. Idempotent.

Run from the repo root, then project:
    python back/data/sport-modules/soccer/apply-rpc-rc1.1-active.py
    python back/data/sport-modules/soccer/project-rpc-workbook.py
"""

import pathlib
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

HERE = pathlib.Path(__file__).parent
WORKBOOK = HERE / "rpc-workbook.rc1.1.xlsm"
HEADER_ROW = 1


def fail(message):
    raise SystemExit(f"STOPPED — {message}")


def rows(ws):
    headers = [c.value for c in ws[HEADER_ROW]]
    col = {str(h): i + 1 for i, h in enumerate(headers) if h}
    out = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        values = {h: ws.cell(row=r, column=i).value for h, i in col.items()}
        if all(v is None or str(v).strip() == "" for v in values.values()):
            continue
        out.append((r, values))
    return col, out


def main():
    wb = openpyxl.load_workbook(WORKBOOK, keep_vba=True)

    _, staging = rows(wb["Implementation Staging"])
    unresolved = [v["mapping_id"] for _, v in staging if v["mapping_status"] == "NEEDS_CANONICAL_ID"]
    if unresolved:
        fail(f"{len(unresolved)} staging rows still need a canonical id: {unresolved[:5]}")

    registry_col, registry = rows(wb["Registry"])
    contexts = [v["rpc_id"] for _, v in registry]
    _, relationships = rows(wb["Relationships"])
    _, properties = rows(wb["Properties"])
    for rpc_id in contexts:
        events = [v for _, v in relationships
                  if v["rpc_id"] == rpc_id and v["related_library"] == "SCORING_EVENT" and v["status"] == "ACTIVE"]
        if not events:
            fail(f"{rpc_id} has no SCORING_EVENT relationship.")
        conditions = [v for _, v in properties if v["rpc_id"] == rpc_id and v["property_category"] == "PRIMARY_SCORING_CONDITION"]
        if len(conditions) != 1:
            fail(f"{rpc_id} has {len(conditions)} PRIMARY_SCORING_CONDITION properties.")

    changed = 0
    for r, values in registry:
        if values["runtime_status"] != "ACTIVE":
            wb["Registry"].cell(row=r, column=registry_col["runtime_status"]).value = "ACTIVE"
            changed += 1

    meta_col, metadata = rows(wb["Metadata"])
    runtime = [(r, v) for r, v in metadata if v["metadata_key"] == "runtime_status"]
    if len(runtime) != 1:
        fail("Metadata must hold exactly one runtime_status row.")
    r, values = runtime[0]
    if values["metadata_value"] != "ACTIVE":
        wb["Metadata"].cell(row=r, column=meta_col["metadata_value"]).value = "ACTIVE"
        changed += 1

    wb.save(WORKBOOK)
    print(f"RPC RC1.1 runtime status ACTIVE ({changed} cells changed; {len(contexts)} contexts).")


if __name__ == "__main__":
    main()
