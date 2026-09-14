"""
Record the scope of the scoring-event vocabulary in the RPC workbook Metadata (Christian, 2026-09-13).

His words: "For RC1.1, I'd treat these as a controlled scoring-event vocabulary used by the Invasion
archetype, without yet making a broader architectural ownership claim." Until now that decision lived
only in a script comment; the workbook said nothing either way. This writes it where the workbook
records its other contracts, as one Metadata row.

Idempotent: re-running leaves an existing, identical row alone and refuses a conflicting one.

Run from the repo root, then project:
    python back/data/sport-modules/soccer/apply-rpc-scoring-vocabulary-scope.py
    python back/data/sport-modules/soccer/project-rpc-workbook.py
"""

import pathlib
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

HERE = pathlib.Path(__file__).parent
WORKBOOK = HERE / "rpc-workbook.rc1.1.xlsm"
HEADER_ROW = 1

ROW = {
    "metadata_key": "scoring_event_vocabulary_scope",
    "metadata_value": "Used with GA-001 Invasion; no broader ownership claim",
    "value_type": "TEXT",
    "required": True,
    "description": (
        "The controlled scoring_event vocabulary is used with GA-001 Invasion for RC1.1, without an "
        "architectural ownership decision. Several events (line crossed, zone entered, gate, target "
        "player) may prove reusable beyond invasion games; ownership is formalized only if later sport "
        "work shows they are archetype-specific. Christian, 2026-09-13: Primary Scoring Event approval."
    ),
    "validation_rule": "Must not assign ownership of the scoring_event vocabulary to a Game Archetype.",
    "introduced_version": "RC1.1",
}


def fail(message):
    raise SystemExit(f"STOPPED — {message}")


def main():
    wb = openpyxl.load_workbook(WORKBOOK, keep_vba=True)
    ws = wb["Metadata"]
    headers = [c.value for c in ws[HEADER_ROW]]
    col = {str(h): i + 1 for i, h in enumerate(headers) if h}
    for key in ROW:
        if key not in col:
            fail(f"Metadata has no '{key}' column.")

    first_empty = None
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        key = ws.cell(row=r, column=col["metadata_key"]).value
        if key is None or str(key).strip() == "":
            if first_empty is None:
                first_empty = r
            continue
        if key == ROW["metadata_key"]:
            existing = {k: ws.cell(row=r, column=col[k]).value for k in ROW}
            if existing != ROW:
                fail(f"a different '{key}' row already exists at row {r}: {existing}")
            print(f"Unchanged: '{key}' already recorded at row {r}.")
            return

    if first_empty is None:
        first_empty = ws.max_row + 1
    for key, value in ROW.items():
        ws.cell(row=first_empty, column=col[key]).value = value
    wb.save(WORKBOOK)
    print(f"Recorded '{ROW['metadata_key']}' at Metadata row {first_empty}.")


if __name__ == "__main__":
    main()
