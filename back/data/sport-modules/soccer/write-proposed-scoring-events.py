"""
Write the PROPOSED Context -> primary scoring event rows for Christian's review.

Christian, 2026-09-13, agreed the responsibility chain

    Realized Game Form -> physically available scoring events -> RPC / Primary Scoring Identity
    -> selected primary scoring event -> activity generation -> coach-facing How to Score

and asked for the first pass of the eight Context -> valid primary scoring event rows from
implementation evidence, "and then for me to review the football/knowledge semantics against each
RPC's Primary Scoring Identity."

THIS IS A DRAFT FOR REVIEW, NOT CANONICAL KNOWLEDGE. Nothing reads the workbook this writes. The
football judgements in the Context sheet are ours to propose and his to decide; the evidence sheets
are what the implementation and real generation actually show.

Evidence sources, all reproducible:
  * game-form setup guidance, and the realizations that add scoring objects:
    back/src/system/sport-module/soccer-module.rc1-v3.json
  * the Primary Scoring Identity of each context: back/src/system/sport-module/rpc-library.rc1.json
  * nine real activities generated 2026-09-13 with src/scripts/run-coach-view-audit.ts (inputs chosen
    from run-archetype-distribution.ts so each context game form is represented)

Run from the repo root:
    python back/data/sport-modules/soccer/write-proposed-scoring-events.py
"""

import json
import pathlib
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8")

HERE = pathlib.Path(__file__).parent
BACK = pathlib.Path(__file__).parents[3]
OUT = HERE / "primary-scoring-events.rc1.1-PROPOSED.xlsx"
RPC_JSON = BACK / "src/system/sport-module/rpc-library.rc1.json"
MODULE_JSON = BACK / "src/system/sport-module/soccer-module.rc1-v3.json"

# ---- 1. the controlled vocabulary he agreed, plus one proposed addition ------------------------------

EVENTS = [
    ("SE-01", "Goal", "The ball crosses the goal line inside a goal.",
     "Goals (full-size, scaled or small); a goalkeeper where the context requires one.", "AGREED"),
    ("SE-02", "Target player", "A pass is received and controlled by a designated target player.",
     "Target player(s) in contrasting bibs, optionally inside a marked target area.", "AGREED"),
    ("SE-03", "Line crossed", "A player dribbles the ball over a marked line, or receives and controls it beyond the line.",
     "A marked line: end line, halfway line or zone line.", "AGREED"),
    ("SE-04", "Target zone entered", "A player dribbles into a marked zone, or receives and controls the ball inside it.",
     "A marked zone.", "AGREED"),
    ("SE-05", "Gate", "The ball is passed or dribbled through a pair of cones or poles.",
     "Gates.", "AGREED"),
    ("SE-06", "Regain", "The team without the ball wins it.",
     "Whatever makes the stated condition observable: a count or timer, a marked zone or line.", "AGREED (as 'regain under a stated condition')"),
    ("SE-07", "Held", "The team without the ball keeps the opponent from crossing a marked line or entering a marked zone for a stated time or number of passes.",
     "A marked line or zone, and a count or timer.",
     "PROPOSED ADDITION — see Counter-Press and Attack Prevention. Already authored in the Delay Reward "
     "realization and in Recover & Reorganize Games (\"defensive success = deny penetration / protect "
     "that space / delay until organized\")."),
]

CONDITION_NOTE = (
    "Refinement from drafting the rows: a stated condition is needed on most events, not only on regain "
    "(Counterattack scores a goal only inside a window after the regain; Goalkeeper Build-Out counts only "
    "from a goalkeeper or controlled start). So the draft treats 'stated condition' as an attribute any "
    "selected event can carry, and 'regain under a stated condition' becomes Regain plus its condition."
)

# ---- 2. game form evidence ---------------------------------------------------------------------------

ADDABLE = (
    "Goals: Goalkeeper Included Condition, Multi-Goal Read. Targets or gates: Variable Target Condition. "
    "Regain in a window: Counter-Press Window. Zones: Zone Structure Condition."
)

GAME_FORMS = {
    "GF1": ("Target zones at both ends.", "SE-04; SE-03", "Not generated in this pass.", "SE-04; SE-03", ""),
    "GF2": ("A directional target line or zone at one end.", "SE-03; SE-04",
            "2 activities: both setups marked end zones; both scored 'a pass or run that beats a defensive line', "
            "not reaching the end zone.", "SE-03; SE-04", ""),
    "GF3": ("A positional grid (typically 3 zones x 3 channels); no goal or target.", "SE-04; SE-03",
            "1 activity ('Defend deeper and stay compact.'): setup marked halves and 'aim to progress into the "
            "opponent's half'; scored 'winning the ball back'. A second intention ('Create better support angles "
            "under pressure.') failed assembly on 3 of 3 attempts.", "SE-04; SE-03",
            "No goals by default, which matters for Finishing realized through Positional Play."),
    "GF4": ("A two-goal or two-target field with a midline marker; the transition moment is any possession change.",
            "SE-01 (goal variant) or SE-02/SE-04 (target variant); SE-03 (midline); SE-06",
            "2 activities ('Counter-attack quickly after winning the ball.'): one setup had goals at each end, one "
            "marked no scoring object at all; BOTH scored 'winning the ball back' — the counter-press side of the "
            "exchange, for a counter-attacking request.", "SE-01; SE-03; SE-06",
            "'Two-target' is not specific: the authored setup should say which target."),
    "GF5": ("Mark the contest area; numerical asymmetry.", "None authored", "Not generated in this pass.", "(none)",
            "Gap: no scoring object is authored."),
    "GF6": ("Target player(s) at one or both ends.", "SE-02", "Not generated in this pass.", "SE-02", ""),
    "GF7": ("3 or 5 lengthwise channels.", "None authored — channels are lanes, not scoring objects",
            "2 activities ('Use the width of the pitch when we attack.'): one setup added an end zone; one marked "
            "only channels while its restart mentions 'after each goal'. Both scored 'a pass or run that beats a "
            "defensive line'.", "(none until an end object is authored)",
            "Gap: Channel Games needs an authored end object (line, zone or goal) before any context can score on it."),
    "GF8": ("Pressing third and defending third with a marked midline; win the ball within the pressing area.",
            "SE-06 (inside the pressing area); SE-03 (midline)", "Not generated in this pass.", "SE-06; SE-03", ""),
    "GF9": ("Full-size or scaled goals at both ends, goalkeepers active.", "SE-01",
            "2 activities ('Finish more of the chances we create.'): both setups had goals and goalkeepers; both "
            "scored 'attacking the open space before the defence recovers'. No goal earned a point.", "SE-01", ""),
    "GF10": ("Open field; the selected constraints impose the structure.", "None authored", "Not generated in this pass.",
             "(depends on constraints)", "Gap: no scoring object is authored."),
    "GF11": ("A recovery window (6-10 s) or marked recovery line; a marked dangerous space to protect; the attack "
             "aims at a real goal or target.", "SE-01/SE-02 (attack); SE-03 (recovery line); SE-04 (protected space); SE-07",
             "Not generated in this pass.", "SE-01; SE-03; SE-04; SE-07", ""),
}

# ---- 3. the eight context rows (the draft he asked for) ------------------------------------------------

CONTEXTS = {
    "RPC-001": (
        "SE-03 Line crossed; SE-04 Target zone entered; SE-05 Gate; SE-02 Target player — each placed beyond the "
        "opponent's first line of pressure",
        "The attack starts from the goalkeeper or a controlled restart, and the ball is under control on arrival. A "
        "long clearance that lands beyond the line does not count (Degenerate Solution: immediate long clearances).",
        "SE-01 Goal rewards finishing, not the establishment of organized attack. SE-06 Regain is the defending side's event.",
        "Is 'under control on arrival' the right observable test for 'established'?"),
    "RPC-002": (
        "SE-03 Line crossed; SE-04 Target zone entered; SE-05 Gate; SE-02 Target player — each placed beyond the "
        "pressing line",
        "The point starts under established pressure, and the team keeps the ball with its next action after "
        "crossing. Escaping and immediately losing it does not count (attacking continuity).",
        "SE-01 Goal is finishing. SE-06 Regain is the press's event — it can be the pressing team's secondary incentive.",
        "Same events as Goalkeeper Build-Out: the difference lives in the starting condition, which matches the "
        "identity rules. Confirm."),
    "RPC-003": (
        "SE-04 Target zone entered (a more advanced zone); SE-03 Line crossed (a marked line between zones); SE-05 Gate; "
        "SE-02 Target player (between lines)",
        "Counts only when the ball arrives in a more advanced area than the one the attack started from. Circulating "
        "possession never scores.",
        "SE-01 Goal belongs to Chance Creation and Finishing. SE-06 Regain is the defending side's event.",
        ""),
    "RPC-004": (
        "SE-04 Target zone entered (a marked finishing zone); SE-05 Gate (a final pass through a gate into the "
        "finishing zone); SE-01 Goal (counts, but is not the only way to score)",
        "The ball is received under control inside the finishing zone with the defence still live. A hopeful ball "
        "into the zone does not count.",
        "SE-06 Regain. SE-03 Line crossed on its own is progression, which Attack Development owns.",
        "A created chance is the hardest of the eight to observe, and zone entry is the closest of the agreed events. "
        "Would 'shot on target' be the better event here?"),
    "RPC-005": (
        "SE-01 Goal",
        "Goalkeeper active wherever the realization includes one (Identity rule: goal, goalkeeper and live defensive "
        "consequence remain representative).",
        "Every other event as PRIMARY. Others may appear only as secondary incentives that keep goals the outcome "
        "(e.g. a finished rebound counts double).",
        "Positional Play has no goals by default, so Finishing through Positional Play needs goals added."),
    "RPC-006": (
        "SE-01 Goal; SE-03 Line crossed; SE-04 Target zone entered; SE-02 Target player — each only inside a stated "
        "window after winning the ball",
        "Scores only inside a stated window after the regain — e.g. 6-10 seconds, the window already authored for "
        "Recover & Reorganize. After the window the same event scores nothing.",
        "SE-06 Regain starts the window but is Counter-Press's event; rewarding it is exactly what both generated "
        "counter-attack activities did.",
        "Window length: seconds or passes, and how long?"),
    "RPC-007": (
        "SE-06 Regain — inside a stated window after losing the ball; SE-07 Held (proposed) — the opponent kept from "
        "crossing the midline inside that window",
        "The window starts the moment the ball is lost; typically 5 seconds (the authored Counter-Press Window).",
        "SE-01, SE-02, SE-03, SE-04 for the pressing team: once the ball is won, attacking it is Counterattack's event.",
        "The identity says 'rather than possession regain alone'. The agreed six can only express the regain. Add "
        "SE-07 Held, so a delay that kills the transition also scores?"),
    "RPC-008": (
        "SE-06 Regain — before the opponent enters the protected zone or crosses the defending line; SE-07 Held "
        "(proposed) — the opponent kept out of the protected zone (typically central, in front of goal) for a stated time",
        "The protected zone is marked; 'lower-value attacking outcomes' are expressed by zone values as a secondary "
        "incentive (Final Third Value, zone weighting), not as the primary event.",
        "SE-01 Goal for the defending team: a goal after winning the ball belongs to Counterattack.",
        "Same SE-07 question as Counter-Press. Confirm zone value stays secondary."),
}

PREFERENCE = {
    "RPC-001": ["SE-03", "SE-04", "SE-05", "SE-02"],
    "RPC-002": ["SE-03", "SE-04", "SE-05", "SE-02"],
    "RPC-003": ["SE-04", "SE-03", "SE-05", "SE-02"],
    "RPC-004": ["SE-04", "SE-05", "SE-01"],
    "RPC-005": ["SE-01"],
    "RPC-006": ["SE-01", "SE-03", "SE-04", "SE-02"],
    "RPC-007": ["SE-06", "SE-07"],
    "RPC-008": ["SE-06", "SE-07"],
}

# What each compatible pair resolves to, and why. Status is one of: RESOLVES, RESOLVES WITH A REALIZATION
# CHOICE, NEEDS AN OBJECT ADDED.
PAIRS = {
    ("RPC-001", "GF2"): ("SE-03", "RESOLVES", "The directional target line."),
    ("RPC-001", "GF3"): ("SE-03", "RESOLVES", "The line into the next grid zone."),
    ("RPC-001", "GF7"): ("", "NEEDS AN OBJECT ADDED", "Channel Games authors no end object."),
    ("RPC-002", "GF2"): ("SE-03", "RESOLVES", "The directional target line, beyond the press."),
    ("RPC-002", "GF3"): ("SE-03", "RESOLVES", "The line into the next grid zone, beyond the press."),
    ("RPC-002", "GF7"): ("", "NEEDS AN OBJECT ADDED", "Channel Games authors no end object."),
    ("RPC-003", "GF3"): ("SE-04", "RESOLVES", "The more advanced grid zone."),
    ("RPC-003", "GF2"): ("SE-04", "RESOLVES", "The target zone."),
    ("RPC-003", "GF7"): ("", "NEEDS AN OBJECT ADDED", "Channel Games authors no end object."),
    ("RPC-004", "GF3"): ("SE-04", "RESOLVES", "The attacking grid zone as the finishing zone."),
    ("RPC-004", "GF2"): ("SE-04", "RESOLVES", "The target zone as the finishing zone."),
    ("RPC-004", "GF7"): ("", "NEEDS AN OBJECT ADDED", "Channel Games authors no end object."),
    ("RPC-005", "GF9"): ("SE-01", "RESOLVES", "Goals and goalkeepers are authored."),
    ("RPC-005", "GF4"): ("SE-01", "RESOLVES WITH A REALIZATION CHOICE", "Only the two-goal variant of Transition Games has goals."),
    ("RPC-005", "GF3"): ("", "NEEDS AN OBJECT ADDED", "Positional Play has no goals; add them (e.g. Goalkeeper Included Condition)."),
    ("RPC-006", "GF4"): ("SE-01", "RESOLVES", "Goal variant; SE-03 on the midline or target otherwise."),
    ("RPC-006", "GF2"): ("SE-03", "RESOLVES", "The directional target line, inside the window."),
    ("RPC-007", "GF4"): ("SE-06", "RESOLVES", "Possession change is the authored transition moment; SE-07 would use the midline."),
    ("RPC-007", "GF2"): ("SE-06", "RESOLVES", "Regain inside the window."),
    ("RPC-008", "GF3"): ("SE-06", "RESOLVES", "Grid zones give the protected zone."),
    ("RPC-008", "GF2"): ("SE-06", "RESOLVES", "Regain before the target line."),
    ("RPC-008", "GF7"): ("SE-06", "RESOLVES", "The central channel as the protected zone."),
}

# ---- 4. the generated activities, as a coach saw them ---------------------------------------------------

GENERATED = [
    ("Keep the ball better under pressure.", "Directional Possession Games", 1, "An end zone",
     "Earn a point for a pass or run that beats a defensive line.", "No"),
    ("Keep the ball better under pressure.", "Directional Possession Games", 3, "Two end zones",
     "Earn a point for a pass or run that beats a defensive line.", "No"),
    ("Counter-attack quickly after winning the ball.", "Transition Games", 1, "None (wide channels and neutral players only)",
     "Earn a point for winning the ball back.", "No object to score; scores the counter-press side"),
    ("Counter-attack quickly after winning the ball.", "Transition Games", 2, "Goals at each end",
     "Earn a point for winning the ball back.", "No; scores the counter-press side"),
    ("Use the width of the pitch when we attack.", "Channel Games", 1, "An end zone behind the opposing team",
     "Earn a point for a pass or run that beats a defensive line.", "No"),
    ("Use the width of the pitch when we attack.", "Channel Games", 2, "None (channels only; restart 'after each goal')",
     "Earn a point for a pass or run that beats a defensive line.", "No object to score"),
    ("Finish more of the chances we create.", "Finishing Games", 1, "Goals and goalkeepers",
     "Earn a point for attacking the open space before the defence recovers.", "No"),
    ("Finish more of the chances we create.", "Finishing Games", 2, "Goals and goalkeepers",
     "Earn a point for attacking the open space before the defence recovers.", "No"),
    ("Defend deeper and stay compact.", "Positional Play Games", 1, "Halves ('progress into the opponent's half')",
     "Earn a point for winning the ball back.", "No (a regain is at least on-theme for a defending request)"),
    ("Create better support angles under pressure.", "Positional Play Games", "-", "-",
     "Assembly failed on 3 of 3 attempts: 'Activity 2 missing skeleton mechanic: Opponent consequence emphasis'.", "-"),
]

README = [
    ("What this is", "A DRAFT for review: which observable events may be the primary scoring event for each "
                     "Representative Performance Context, drafted from implementation evidence. Not canonical; nothing reads it."),
    ("The agreed chain", "Realized Game Form -> physically available scoring events -> RPC / Primary Scoring Identity -> "
                         "selected primary scoring event -> activity generation -> coach-facing How to Score."),
    ("Why it matters (evidence)", "Nine real activities generated on 2026-09-13. Seven setups marked a scoring object (end "
                                  "zones, goals with goalkeepers, halves); none of the seven scored it. Both finishing "
                                  "games had goals and goalkeepers and awarded points for 'attacking the open space'. "
                                  "Both counter-attack games awarded points for winning the ball back. See 'Generated Evidence'."),
    ("How to read", "Event Vocabulary: the six agreed events plus one proposed. Game Form Evidence: what each game form "
                    "physically offers. Context Primary Events: the eight rows to review. Context x Game Form: whether "
                    "each compatible pair resolves to an event."),
    ("Stated condition", CONDITION_NOTE),
    ("Decisions requested", "1. The football semantics of each Context row. 2. Add SE-07 Held (Counter-Press, Attack "
                            "Prevention)? 3. Chance Creation: zone entry, or 'shot on target'? 4. Window lengths "
                            "(Counterattack, Counter-Press). 5. Channel Games needs an authored end object."),
    ("Where the rows would live", "Without a schema change: game form -> available events in the Sport Module's existing "
                                  "scoring_structure_type column (present on all 11 game forms, currently empty and read by "
                                  "nothing); context -> valid events as Relationships (related_library SCORING_EVENT) with the "
                                  "condition as a Property. Neither related_library nor property_category is a controlled "
                                  "vocabulary column in the frozen schema. The event vocabulary itself needs one owner."),
]

HEADER_FILL = PatternFill("solid", fgColor="DDE6F0")


def sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


def main():
    rpc = json.loads(RPC_JSON.read_text(encoding="utf-8"))
    module = json.loads(MODULE_JSON.read_text(encoding="utf-8"))
    name_of_context = {r["rpc_id"]: r["rpc_name"] for r in rpc["registry"]}
    scoring_identity = {s["rpc_id"]: s["statement_text"] for s in rpc["statements"] if s["statement_type"] == "PRIMARY_SCORING_IDENTITY"}
    form_name = {g["game_form_id"]: g["game_form_name"] for g in module["game_forms"]}
    event_name = {e[0]: e[1] for e in EVENTS}

    # Every compatible pair in the canonical workbook must be drafted, and nothing else — so the draft
    # cannot silently skip a pair or invent one.
    compatible = {(r["rpc_id"], r["related_id"]): r["relationship_strength"] for r in rpc["relationships"]
                  if r["related_library"] == "GAME_FORM" and r["status"] == "ACTIVE"}
    if set(compatible) != set(PAIRS):
        raise SystemExit(f"Draft pairs differ from the workbook: missing {sorted(set(compatible) - set(PAIRS))}, "
                         f"extra {sorted(set(PAIRS) - set(compatible))}")
    if set(CONTEXTS) != set(name_of_context) or set(GAME_FORMS) != set(form_name):
        raise SystemExit("Draft contexts or game forms differ from the canonical libraries.")
    for (rpc_id, _), (event, status, _) in PAIRS.items():
        if event and event not in PREFERENCE[rpc_id]:
            raise SystemExit(f"{rpc_id} resolves to {event}, which is not one of its valid events.")
        if bool(event) == status.startswith("NEEDS"):
            raise SystemExit(f"{rpc_id}: status '{status}' disagrees with event '{event}'.")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet(wb, "README", ["Topic", "Detail"], README, [26, 120])
    sheet(wb, "Event Vocabulary", ["event_id", "Event", "Observable definition (what a coach counts)", "Needs marked out", "Status"],
          EVENTS, [10, 20, 60, 45, 60])
    sheet(wb, "Game Form Evidence",
          ["game_form_id", "Game form", "Authored setup (what is marked out)", "Events available by default",
           "Events a selected realization can add", "Observed in generated activities (2026-09-13)",
           "Proposed scoring_structure_type", "Note"],
          [(gf, form_name[gf], s, avail, ADDABLE, observed, proposed, note)
           for gf, (s, avail, observed, proposed, note) in GAME_FORMS.items()],
          [10, 24, 40, 30, 36, 60, 22, 36])
    sheet(wb, "Context Primary Events",
          ["rpc_id", "Context", "Primary Scoring Identity (canonical, unchanged)", "Valid primary scoring events (in preference order)",
           "Stated condition (draft)", "Not valid as primary, and why", "Question for review"],
          [(rid, name_of_context[rid], scoring_identity[rid], *CONTEXTS[rid]) for rid in sorted(CONTEXTS)],
          [10, 20, 40, 50, 50, 45, 45])
    sheet(wb, "Context x Game Form",
          ["rpc_id", "Context", "game_form_id", "Game form", "Compatibility", "Selected primary event", "Status", "Why"],
          [(rid, name_of_context[rid], gf, form_name[gf], compatible[(rid, gf)],
            f"{ev} {event_name[ev]}" if ev else "(none yet)", status, why)
           for (rid, gf), (ev, status, why) in sorted(PAIRS.items(), key=lambda kv: (kv[0][0], ["PRIMARY", "SECONDARY"].index(compatible[kv[0]])))],
          [10, 20, 12, 26, 14, 24, 30, 50])
    sheet(wb, "Generated Evidence",
          ["Coaching intention (input)", "Game form selected", "Slot", "Scoring object in the setup", "What Scoring rewarded", "Scored the object?"],
          GENERATED, [36, 26, 6, 36, 55, 34])

    wb.save(OUT)
    counts = {}
    for _, status, _ in PAIRS.values():
        counts[status] = counts.get(status, 0) + 1
    print(f"Wrote {OUT.name}: {len(EVENTS)} events, {len(GAME_FORMS)} game forms, {len(CONTEXTS)} contexts, "
          f"{len(PAIRS)} pairs {counts}, {len(GENERATED)} generated activities.")


if __name__ == "__main__":
    main()
