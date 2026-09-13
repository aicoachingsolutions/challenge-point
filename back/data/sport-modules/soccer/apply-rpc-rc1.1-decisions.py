"""
Apply Christian's RC1.1 follow-up decisions (2026-09-13) to the RPC workbook.

1. AFFORDANCE TARGETS DEFERRED. The 43 Affordance Target staging rows move from NEEDS_CANONICAL_ID to
   DEFERRED — "Mapping intentionally postponed", in the workbook's own controlled vocabulary. His
   words: "I do not want canonical IDs inferred simply to get the RPC Library into runtime." DEFERRED
   records that the mappings remain unresolved: it creates no relationship and infers no id.

   With them deferred, no NEEDS_CANONICAL_ID row remains, so the Workbook Standard no longer blocks
   ACTIVE. runtime_status is deliberately NOT changed here: he ties ACTIVE to reviewing the
   Context -> primary scoring event rows, so that switch is his decision and a separate edit.

2. STALE AUDIT ROW REMOVED. The RC1.1 Audit sheet recorded "Implementation Staging OLD retained as
   archive", but no such sheet exists. His instruction: remove the stale reference rather than
   restore an obsolete sheet.

Idempotent. Run from the repo root, before resolution (which leaves DEFERRED rows alone) and
projection:
    python back/data/sport-modules/soccer/apply-rpc-rc1.1-decisions.py
    python back/data/sport-modules/soccer/resolve-rpc-staging.py
    python back/data/sport-modules/soccer/project-rpc-workbook.py
"""

import pathlib
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

HERE = pathlib.Path(__file__).parent
WORKBOOK = HERE / "rpc-workbook.rc1.1.xlsm"
HEADER_ROW = 1

EXPECTED_AFFORDANCE_ROWS = 43
DEFERRAL_NOTE = (
    "Deferred by Christian, 2026-09-13: canonical Affordance Target ids are not to be inferred to reach "
    "runtime; revisit with the Affordance Target relationship work."
)
STALE_AUDIT_TEXT = "Implementation Staging OLD"


def fail(message):
    raise SystemExit(f"STOPPED — {message}")


def main():
    wb = openpyxl.load_workbook(WORKBOOK, keep_vba=True)

    staging = wb["Implementation Staging"]
    headers = [c.value for c in staging[HEADER_ROW]]
    col = {str(h): i + 1 for i, h in enumerate(headers) if h}
    deferred, already = 0, 0
    for r in range(HEADER_ROW + 1, staging.max_row + 1):
        if staging.cell(row=r, column=col["target_library"]).value != "AFFORDANCE":
            continue
        status_cell = staging.cell(row=r, column=col["mapping_status"])
        notes_cell = staging.cell(row=r, column=col["notes"])
        mapping_id = staging.cell(row=r, column=col["mapping_id"]).value
        if status_cell.value == "DEFERRED" and DEFERRAL_NOTE in str(notes_cell.value or ""):
            already += 1
            continue
        if status_cell.value != "NEEDS_CANONICAL_ID":
            fail(f"{mapping_id} is {status_cell.value}; only unresolved Affordance Target rows are deferred.")
        status_cell.value = "DEFERRED"
        existing = str(notes_cell.value or "").strip()
        notes_cell.value = f"{DEFERRAL_NOTE} | {existing}" if existing else DEFERRAL_NOTE
        deferred += 1
    if deferred + already != EXPECTED_AFFORDANCE_ROWS:
        fail(f"found {deferred + already} Affordance Target rows; the decision covers {EXPECTED_AFFORDANCE_ROWS}.")

    audit = wb["RC1.1 Audit"]
    stale_rows = [
        r for r in range(HEADER_ROW + 1, audit.max_row + 1)
        if any(STALE_AUDIT_TEXT in str(audit.cell(row=r, column=c).value or "") for c in range(1, audit.max_column + 1))
    ]
    if len(stale_rows) > 1:
        fail(f"RC1.1 Audit mentions '{STALE_AUDIT_TEXT}' on {len(stale_rows)} rows; expected one.")
    for r in stale_rows:
        audit.delete_rows(r, 1)

    wb.save(WORKBOOK)
    print(f"Affordance Target rows deferred: {deferred}; already deferred: {already}.")
    print(f"Stale audit rows removed: {len(stale_rows)}.")


if __name__ == "__main__":
    main()
