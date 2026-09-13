"""
Apply Christian's Primary Scoring Event approval (2026-09-13) to the RPC workbook.

His decisions, as approved:
  * A primary scoring event is an OBSERVABLE EVENT plus a QUALIFYING CONDITION; conditions attach to
    any event.
  * Seven controlled events: goal, target player, line crossed, target zone entered, gate, regain,
    denial. "Denial" replaces the proposed "Held"; no shot-on-target event.
  * The eight Context rows, with the valid events in the order he approved and the condition he set
    for each: the existing 5-second window for Counter-Press, and no fixed time for Attack Prevention.
  * The vocabulary is used with GA-001 Invasion for now, with no broader ownership claim.

Where each lands, using only sheets the frozen schema already has (no new sheet or column):
  Controlled Vocabulary  vocabulary "scoring_event": one row per event, its observable definition
                         as the description.
  Relationships          related_library SCORING_EVENT: one row per valid event per context, in
                         his order (relationship ids preserve it).
  Properties             property_category PRIMARY_SCORING_CONDITION: one per context, his wording.

Runtime status is NOT changed here. He made ACTIVE conditional on these changes being reflected
AND validation passing, so ACTIVE is a separate, later edit.

Idempotent: a row that already exists with the same content is left alone; one that exists with
different content stops the script, because that would be a different decision.

Run from the repo root, then resolve (no-op for these rows) and project:
    python back/data/sport-modules/soccer/apply-rpc-scoring-events.py
    python back/data/sport-modules/soccer/resolve-rpc-staging.py
    python back/data/sport-modules/soccer/project-rpc-workbook.py
"""

import pathlib
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

HERE = pathlib.Path(__file__).parent
WORKBOOK = HERE / "rpc-workbook.rc1.1.xlsm"
HEADER_ROW = 1
PROVENANCE = "Christian, 2026-09-13: Primary Scoring Event approval (RC1.1)"

EVENTS = [
    ("goal", "The ball crosses the goal line inside a goal."),
    ("target_player", "A pass is received and controlled by a designated target player."),
    ("line_crossed", "A player dribbles the ball over a marked line, or receives and controls it beyond the line."),
    ("target_zone_entered", "A player dribbles into a marked zone, or receives and controls the ball inside it."),
    ("gate", "The ball is passed or dribbled through a pair of cones or poles."),
    ("regain", "The team without the ball wins it."),
    ("denial", "A specified opponent progression event does not occur before the stated defensive condition is satisfied."),
]

CONTEXT_EVENTS = {
    "RPC-001": ["line_crossed", "target_zone_entered", "gate", "target_player"],
    "RPC-002": ["line_crossed", "target_zone_entered", "gate", "target_player"],
    "RPC-003": ["target_zone_entered", "line_crossed", "gate", "target_player"],
    "RPC-004": ["target_zone_entered", "gate", "target_player"],
    "RPC-005": ["goal"],
    "RPC-006": ["goal", "line_crossed", "target_zone_entered", "target_player"],
    "RPC-007": ["regain", "denial"],
    "RPC-008": ["regain", "denial"],
}

CONDITIONS = {
    "RPC-001": "Initiated from the goalkeeper or a controlled defensive restart, progressing beyond the initial "
               "pressure or progression line, with possession controlled on arrival. A long clearance that "
               "happens to reach the target does not satisfy the event.",
    "RPC-002": "The sequence begins under established coordinated pressure, the event occurs beyond the pressing "
               "line, and the attacking team retains possession through the next action.",
    "RPC-003": "Purposeful advancement from the attack's starting condition rather than possession alone: reaching "
               "a more advanced representative area while maintaining attacking possession.",
    "RPC-004": "The finishing area is reached through the event and the ball is received under control, with live "
               "defensive opposition. The context ends when the meaningful scoring opportunity has been created.",
    "RPC-005": "Where the realization includes a goalkeeper, the goalkeeper remains active as part of the "
               "representative condition.",
    "RPC-006": "Counts only within the representative transition window following the regain.",
    "RPC-007": "Within the existing 5-second Counter-Press Window following possession loss. Denial: the opponent "
               "does not cross the specified progression line or enter the specified progression zone within "
               "that window.",
    "RPC-008": "The opponent is prevented from crossing the specified progression line, entering the protected zone, "
               "or reaching the realized attacking objective during the attacking sequence. No fixed time is "
               "intrinsic to this context; a particular realization may use time where appropriate.",
}


def fail(message):
    raise SystemExit(f"STOPPED — {message}")


def table(ws):
    headers = [c.value for c in ws[HEADER_ROW]]
    col = {str(h): i + 1 for i, h in enumerate(headers) if h}
    rows = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        values = {h: ws.cell(row=r, column=i).value for h, i in col.items()}
        if all(v is None or str(v).strip() == "" for v in values.values()):
            continue
        rows.append((r, values))
    return col, rows


def append(ws, col, values):
    """Write at the first empty row after the table — ws.append would land after 1000 formatted rows."""
    _, rows = table(ws)
    target = (rows[-1][0] + 1) if rows else HEADER_ROW + 1
    for header, value in values.items():
        if header not in col:
            fail(f"sheet '{ws.title}' has no column '{header}'.")
        ws.cell(row=target, column=col[header]).value = value


def next_id(rows, column, prefix):
    numbers = [int(str(v[column]).split("-")[-1]) for _, v in rows if v.get(column)]
    return f"{prefix}{max(numbers) + 1:03d}"


def main():
    wb = openpyxl.load_workbook(WORKBOOK, keep_vba=True)
    counts = {"vocabulary": 0, "relationships": 0, "properties": 0}

    # 1. Controlled vocabulary.
    vocab = wb["Controlled Vocabulary"]
    vcol, vrows = table(vocab)
    existing = {(v["vocabulary"], v["value"]): v["description"] for _, v in vrows}
    for key, definition in EVENTS:
        if ("scoring_event", key) in existing:
            if existing[("scoring_event", key)] != definition:
                fail(f"scoring_event '{key}' exists with a different definition.")
            continue
        append(vocab, vcol, {"vocabulary": "scoring_event", "value": key, "description": definition})
        counts["vocabulary"] += 1

    # 2. Relationships, in his order.
    rel = wb["Relationships"]
    rcol, rrows = table(rel)
    registry = {v["rpc_id"] for _, v in table(wb["Registry"])[1]}
    have = [(v["rpc_id"], v["related_id"]) for _, v in rrows if v["related_library"] == "SCORING_EVENT"]
    for rpc_id, events in CONTEXT_EVENTS.items():
        if rpc_id not in registry:
            fail(f"{rpc_id} is not in the Registry.")
        existing_for_context = [e for r, e in have if r == rpc_id]
        if existing_for_context:
            if existing_for_context != events:
                fail(f"{rpc_id} already has SCORING_EVENT rows {existing_for_context}, not the approved {events}.")
            continue
        for position, event in enumerate(events, start=1):
            _, rrows = table(rel)
            append(rel, rcol, {
                "relationship_id": next_id(rrows, "relationship_id", "RPC-REL-"),
                "rpc_id": rpc_id,
                "related_library": "SCORING_EVENT",
                "related_id": event,
                "relationship_type": "COMPATIBLE",
                "relationship_strength": "PRIMARY",
                "status": "ACTIVE",
                "provenance": PROVENANCE,
                "notes": f"Valid primary scoring event {position} of {len(events)}, in the approved order.",
            })
            counts["relationships"] += 1

    # 3. Qualifying conditions.
    props = wb["Properties"]
    pcol, prows = table(props)
    for rpc_id, condition in CONDITIONS.items():
        current = [v for _, v in prows if v["rpc_id"] == rpc_id and v["property_category"] == "PRIMARY_SCORING_CONDITION"]
        if current:
            if len(current) != 1 or current[0]["property_text"] != condition:
                fail(f"{rpc_id} already has a different PRIMARY_SCORING_CONDITION.")
            continue
        _, prows = table(props)
        append(props, pcol, {
            "property_id": next_id(prows, "property_id", "RPC-PROP-"),
            "rpc_id": rpc_id,
            "property_category": "PRIMARY_SCORING_CONDITION",
            "property_text": condition,
            "importance": "CORE",
            "status": "ACTIVE",
            "provenance": PROVENANCE,
        })
        _, prows = table(props)
        counts["properties"] += 1

    wb.save(WORKBOOK)
    print(f"Added: {counts['vocabulary']} scoring_event vocabulary rows, {counts['relationships']} SCORING_EVENT "
          f"relationships, {counts['properties']} PRIMARY_SCORING_CONDITION properties.")


if __name__ == "__main__":
    main()
